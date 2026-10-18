from datetime import datetime, timedelta
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import os
import smtplib
import time
import numpy as np
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
 
# ==========================================
# 1. CẤU HÌNH DANH MỤC WATCHLIST CHUNG (16 MÃ HOSE)
# ==========================================
WATCHLIST_CONFIG = [
    # Công nghệ & Viễn thông
    {"symbol": "FPT", "sector": "Công nghệ", "fa_base": 88},
    {"symbol": "CMG", "sector": "Công nghệ", "fa_base": 78},
    {"symbol": "ELC", "sector": "Công nghệ", "fa_base": 75},
    # Ngân hàng
    {"symbol": "TCB", "sector": "Ngân hàng", "fa_base": 85},
    {"symbol": "MBB", "sector": "Ngân hàng", "fa_base": 84},
    {"symbol": "ACB", "sector": "Ngân hàng", "fa_base": 82},
    {"symbol": "VPB", "sector": "Ngân hàng", "fa_base": 79},
    {"symbol": "STB", "sector": "Ngân hàng", "fa_base": 78},
    {"symbol": "TPB", "sector": "Ngân hàng", "fa_base": 76},
    {"symbol": "VCB", "sector": "Ngân hàng", "fa_base": 90},
    {"symbol": "BID", "sector": "Ngân hàng", "fa_base": 83},
    {"symbol": "CTG", "sector": "Ngân hàng", "fa_base": 82},
    # Chứng khoán & Khác
    {"symbol": "SSI", "sector": "Chứng khoán", "fa_base": 81},
    {"symbol": "VND", "sector": "Chứng khoán", "fa_base": 74},
    {"symbol": "HPG", "sector": "Thép", "fa_base": 86},
    {"symbol": "PDR", "sector": "Bất động sản", "fa_base": 72},
]
 
# Máy Windows cá nhân: lưu vào D:\HOSE_Code. GitHub Actions (Linux): lưu vào thư mục output/.
# Có thể ghi đè bằng biến môi trường OUTPUT_DIR.
OUTPUT_DIR = os.environ.get("OUTPUT_DIR") or (
    r"D:\HOSE_Code" if os.name == "nt" else "output"
)
 
# Dòng cảnh báo hiển thị đầu email. Đặt thành chuỗi rỗng "" khi đã dùng dữ liệu thị trường thật.
SIMULATED_DATA_NOTICE = (
    "Lưu ý: điểm số trong báo cáo này được sinh ngẫu nhiên (mô phỏng), "
    "chưa lấy từ dữ liệu giá/khối lượng thực tế, không dùng làm cơ sở ra quyết định đầu tư."
)
 
 
# ==========================================
# MODULE 1: ĐÁNH GIÁ TỔNG QUAN WATCHLIST
# ==========================================
def run_main_watchlist_report():
  print("🚀 [1/3] Đang chạy báo cáo Tổng quan Watchlist...")
  np.random.seed(42)
  rows = []
  for item in WATCHLIST_CONFIG:
    sym = item["symbol"]
    sec = item["sector"]
    fa_score = item["fa_base"] + np.random.randint(-2, 3)
    vol_vs_ma20 = round(np.random.uniform(1.2, 2.4), 2)
    price_action = np.random.choice(
        ["Tích lũy siết nền (VCP)", "Breakout khỏi nền giá", "Vượt đỉnh ngắn hạn"]
    )
    macd_signal = np.random.choice(
        ["Cắt lên (Bullish)", "Duy trì Dương", "Phân kỳ dương"]
    )
    tech_score = int(
        min(
            98,
            70
            + (vol_vs_ma20 * 8)
            + (10 if "Breakout" in price_action or "Vượt" in price_action else 5),
        )
    )
    nlp_score = np.random.randint(72, 93)
    total_score = round(
        tech_score * 0.45 + fa_score * 0.35 + nlp_score * 0.20, 1
    )
 
    if total_score >= 82:
      recommendation = "MUA MẠNH (Strong Buy)"
      reliability = "Rất Cao (High Conviction)"
    elif total_score >= 76:
      recommendation = "MUA TÍCH LŨY (Accumulate)"
      reliability = "Thích Hợp (Reliable)"
    else:
      recommendation = "QUAN SÁT (Watch)"
      reliability = "Trung Bình (Moderate)"
 
    rows.append({
        "Mã CK": sym,
        "Nhóm Ngành": sec,
        "Điểm Kỹ Thuật (45%)": tech_score,
        "Điểm Cơ Bản FA (35%)": fa_score,
        "Điểm Sentiment NLP (20%)": nlp_score,
        "Tổng Điểm Đánh Giá": total_score,
        "Khối Lượng vs MA20": f"{vol_vs_ma20}x",
        "Tín Hiệu Kỹ Thuật Nổi Bật": (
            f"{price_action} | MACD: {macd_signal}"
        ),
        "Mức Độ Đáng Tin": reliability,
        "Khuyến Nghị Đánh Giá": recommendation,
    })
 
  df = pd.DataFrame(rows)
  df = df.sort_values(by="Tổng Điểm Đánh Giá", ascending=False).reset_index(
      drop=True
  )
 
  os.makedirs(OUTPUT_DIR, exist_ok=True)
  file_path = os.path.join(OUTPUT_DIR, "HOSE_Watchlist_Report.xlsx")
 
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Đánh Giá Watchlist Daily"
 
  ws.merge_cells("A1:J1")
  title_cell = ws["A1"]
  title_cell.value = (
      "BÁO CÁO ĐÁNH GIÁ CỔ PHIẾU TIỀM NĂNG DAILY (HOSE WATCHLIST)"
  )
  title_cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
  title_cell.fill = PatternFill(
      start_color="1F497D", end_color="1F497D", fill_type="solid"
  )
  title_cell.alignment = Alignment(horizontal="center", vertical="center")
  ws.row_dimensions[1].height = 35
 
  headers = list(df.columns)
  ws.row_dimensions[3].height = 28
  header_fill = PatternFill(
      start_color="244062", end_color="244062", fill_type="solid"
  )
  header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
  border_thin = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )
 
  for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    cell.border = border_thin
 
  for row_idx, row_data in enumerate(df.to_dict("records"), 4):
    ws.row_dimensions[row_idx].height = 22
    for col_idx, key in enumerate(headers, 1):
      val = row_data[key]
      cell = ws.cell(row=row_idx, column=col_idx, value=val)
      cell.font = Font(name="Calibri", size=10)
      cell.border = border_thin
 
      if key in [
          "Mã CK",
          "Nhóm Ngành",
          "Khối Lượng vs MA20",
          "Mức Độ Đáng Tin",
      ]:
        cell.alignment = Alignment(horizontal="center", vertical="center")
      elif key in [
          "Điểm Kỹ Thuật (45%)",
          "Điểm Cơ Bản FA (35%)",
          "Điểm Sentiment NLP (20%)",
          "Tổng Điểm Đánh Giá",
      ]:
        cell.alignment = Alignment(horizontal="right", vertical="center")
      else:
        cell.alignment = Alignment(horizontal="left", vertical="center")
 
      if key == "Khuyến Nghị Đánh Giá":
        if "MUA MẠNH" in str(val):
          cell.fill = PatternFill(
              start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="375623"
          )
        elif "MUA TÍCH LŨY" in str(val):
          cell.fill = PatternFill(
              start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="7F6000"
          )
 
  for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
  ws.column_dimensions["H"].width = 38
 
  chart = BarChart()
  chart.type = "col"
  chart.style = 10
  chart.title = "Xếp Hạng Tổng Điểm Đánh Giá Cổ Phiếu Watchlist"
  chart.y_axis.title = "Điểm Đánh Giá"
  chart.x_axis.title = "Mã Cổ Phiếu"
  data_ref = Reference(ws, min_col=6, min_row=3, max_row=3 + len(df))
  cats_ref = Reference(ws, min_col=1, min_row=4, max_row=3 + len(df))
  chart.add_data(data_ref, titles_from_data=True)
  chart.set_categories(cats_ref)
  chart.width = 18
  chart.height = 12
  ws.add_chart(chart, "L3")
 
  wb.save(file_path)
  print(f"   -> Đã lưu: {file_path}")
  return df
 
 
# ==========================================
# MODULE 2: BỘ LỌC PULLBACK / TEST CUNG
# ==========================================
def run_pullback_analysis():
  print("🔍 [2/3] Đang chạy bộ lọc Pullback / Test cung...")
  np.random.seed(datetime.now().second + datetime.now().minute)
  rows = []
  for item in WATCHLIST_CONFIG:
    sym = item["symbol"]
    sec = item["sector"]
    fa_score = item["fa_base"] + np.random.randint(-2, 3)
    market_action = np.random.choice([
        "Test cung hỗ trợ (Pullback nhẹ, Vol thấp)",
        "Tích lũy siết nền (VCP)",
        "Áp lực chốt lời ngắn hạn (Đỏ nhẹ cuối phiên)",
        "Breakout vượt đỉnh khối lượng lớn",
    ])
 
    if "Test cung" in market_action:
      tech_score = np.random.randint(85, 94)
      vol_desc = "Thấp hơn MA20 (Cạn cung)"
      strategy = "MUA GOM (Tích lũy giá đỏ)"
      reliability = "Rất Cao (Điểm mua an toàn)"
    elif "Tích lũy" in market_action:
      tech_score = np.random.randint(78, 86)
      vol_desc = "Ổn định quanh MA20"
      strategy = "MUA TÍCH LŨY DẦN"
      reliability = "Thích Hợp (Chờ bùng nổ)"
    elif "Đỏ nhẹ" in market_action:
      tech_score = np.random.randint(70, 78)
      vol_desc = "Bình thường (Test MA ngắn hạn)"
      strategy = "QUAN SÁT (Canh nhúng sâu thêm)"
      reliability = "Trung Bình (Chờ xác nhận)"
    else:
      tech_score = np.random.randint(88, 97)
      vol_desc = "Đột biến > 2x MA20"
      strategy = "MUA ĐUỔI THẬN TRỌNG (Tỷ trọng nhỏ)"
      reliability = "Cao (Dòng tiền mạnh)"
 
    nlp_score = np.random.randint(72, 92)
    total_score = round(
        tech_score * 0.45 + fa_score * 0.35 + nlp_score * 0.20, 1
    )
 
    rows.append({
        "Mã CK": sym,
        "Nhóm Ngành": sec,
        "Điểm Kỹ Thuật (45%)": tech_score,
        "Điểm Cơ Bản FA (35%)": fa_score,
        "Điểm Sentiment NLP (20%)": nlp_score,
        "Tổng Điểm Đánh Giá": total_score,
        "Trạng Thái Cung Cầu": market_action,
        "Đánh Giá Thanh Khoản": vol_desc,
        "Mức Độ Đáng Tin": reliability,
        "Chiến Lược Hành Động": strategy,
    })
 
  df = pd.DataFrame(rows)
  df = df.sort_values(by="Tổng Điểm Đánh Giá", ascending=False).reset_index(
      drop=True
  )
 
  file_path = os.path.join(OUTPUT_DIR, "HOSE_Watchlist_Pullback_Report.xlsx")
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Đánh Giá Pullback"
 
  ws.merge_cells("A1:J1")
  title_cell = ws["A1"]
  title_cell.value = (
      "BÁO CÁO ĐÁNH GIÁ WATCHLIST & ĐIỂM MUA AN TOÀN (PULLBACK / TEST CUNG)"
  )
  title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
  title_cell.fill = PatternFill(
      start_color="1F497D", end_color="1F497D", fill_type="solid"
  )
  title_cell.alignment = Alignment(horizontal="center", vertical="center")
  ws.row_dimensions[1].height = 35
 
  headers = list(df.columns)
  ws.row_dimensions[3].height = 28
  header_fill = PatternFill(
      start_color="244062", end_color="244062", fill_type="solid"
  )
  header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
  border_thin = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )
 
  for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    cell.border = border_thin
 
  for row_idx, row_data in enumerate(df.to_dict("records"), 4):
    ws.row_dimensions[row_idx].height = 22
    for col_idx, key in enumerate(headers, 1):
      val = row_data[key]
      cell = ws.cell(row=row_idx, column=col_idx, value=val)
      cell.font = Font(name="Calibri", size=10)
      cell.border = border_thin
 
      if key in ["Mã CK", "Nhóm Ngành", "Mức Độ Đáng Tin"]:
        cell.alignment = Alignment(horizontal="center", vertical="center")
      elif key in [
          "Điểm Kỹ Thuật (45%)",
          "Điểm Cơ Bản FA (35%)",
          "Điểm Sentiment NLP (20%)",
          "Tổng Điểm Đánh Giá",
      ]:
        cell.alignment = Alignment(horizontal="right", vertical="center")
      else:
        cell.alignment = Alignment(horizontal="left", vertical="center")
 
      if key == "Chiến Lược Hành Động":
        cell.font = Font(name="Calibri", size=10, bold=True)
        if "MUA GOM" in str(val):
          cell.fill = PatternFill(
              start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="375623"
          )
        elif "MUA TÍCH LŨY" in str(val):
          cell.fill = PatternFill(
              start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="7F6000"
          )
        elif "QUAN SÁT" in str(val):
          cell.fill = PatternFill(
              start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="C65911"
          )
 
  for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
  ws.column_dimensions["G"].width = 38
  ws.column_dimensions["J"].width = 32
 
  chart = BarChart()
  chart.type = "col"
  chart.style = 10
  chart.title = "Xếp Hạng Điểm Đánh Giá & Điểm Mua Pullback"
  chart.y_axis.title = "Tổng Điểm Đánh Giá"
  chart.x_axis.title = "Mã Cổ Phiếu"
  data_ref = Reference(ws, min_col=6, min_row=3, max_row=3 + len(df))
  cats_ref = Reference(ws, min_col=1, min_row=4, max_row=3 + len(df))
  chart.add_data(data_ref, titles_from_data=True)
  chart.set_categories(cats_ref)
  chart.width = 18
  chart.height = 12
  ws.add_chart(chart, "L3")
 
  wb.save(file_path)
  print(f"   -> Đã lưu: {file_path}")
  return df
 
 
# ==========================================
# MODULE 3: BỘ LỌC TÍN HIỆU BÁN (STOP-LOSS / TAKE PROFIT)
# ==========================================
def run_technical_sell_screener():
  print("📉 [3/3] Đang chạy bộ lọc Tín hiệu Kỹ thuật Bán...")
  np.random.seed(datetime.now().hour + datetime.now().minute)
  rows = []
  for item in WATCHLIST_CONFIG:
    sym = item["symbol"]
    sec = item["sector"]
    rsi = np.random.randint(35, 82)
    macd_status = np.random.choice([
        "Cắt xuống (Bearish Cross)",
        "Duy trì Dương (Bullish)",
        "Phân kỳ âm (Top Divergence)",
        "Suy yếu quanh đường Zero",
    ])
    ma_status = np.random.choice([
        "Thủng MA20 (Gãy xu hướng ngắn)",
        "Giữ trên MA20 (An toàn)",
        "Test MA50 (Hỗ trợ cứng)",
        "Rơi thủng MA50 (Xác nhận downtrend)",
    ])
    vol_action = np.random.choice([
        "Bán tháo khối lượng lớn (Phân phối)",
        "Cạn cung (Vol thấp, đi ngang)",
        "Bình thường",
        "Kéo xả cuối phiên (ATC)",
    ])
    profit_status = np.random.choice([
        "Đạt lãi kỳ vọng > 15%",
        "Hòa vốn / Lãi nhẹ 3-5%",
        "Vi phạm âm 7% (Cắt lỗ)",
        "Đang gồng lời tốt",
    ])
 
    sell_risk_score = 30
    if "Thủng MA20" in ma_status or "Rơi thủng MA50" in ma_status:
      sell_risk_score += 35
    if "Cắt xuống" in macd_status or "Phân kỳ âm" in macd_status:
      sell_risk_score += 20
    if "Bán tháo" in vol_action or "Kéo xả" in vol_action:
      sell_risk_score += 25
    if rsi > 75:
      sell_risk_score += 15
    elif rsi < 40:
      sell_risk_score -= 10
 
    sell_risk_score = int(min(100, max(10, sell_risk_score)))
 
    if (
        "Vi phạm âm 7%" in profit_status
        or sell_risk_score >= 80
        or "Rơi thủng MA50" in ma_status
    ):
      sell_action = "CẮT LỖ / BÁN KHẨN CẤP (Stop-Loss 100%)"
      priority = "Cao nhất (Nguy hiểm)"
    elif sell_risk_score >= 65 or "Đạt lãi kỳ vọng" in profit_status:
      sell_action = "CHỐT LỜI / HẠ TỶ TRỌNG (Take Profit 50-100%)"
      priority = "Cao (Hiện thực hóa lợi nhuận)"
    elif sell_risk_score >= 45:
      sell_action = "THEO DÕI SÁT / CHẶN LÃI (Trailing Stop)"
      priority = "Trung bình (Cảnh báo)"
    else:
      sell_action = "TIẾP TỤC NẮM GIỮ (Hold & Let Profit Run)"
      priority = "Thấp (An toàn)"
 
    rows.append({
        "Mã CK": sym,
        "Nhóm Ngành": sec,
        "Điểm Rủi Ro Bán (0-100)": sell_risk_score,
        "Chỉ báo RSI": rsi,
        "Trạng thái MACD": macd_status,
        "Vị thế MA (MA20/MA50)": ma_status,
        "Khối lượng & Cung Cầu": vol_action,
        "Trạng thái Lợi Nhuận": profit_status,
        "Mức Độ Ưu Tiên": priority,
        "Hành Động Khuyến Nghị": sell_action,
    })
 
  df = pd.DataFrame(rows)
  df = df.sort_values(by="Điểm Rủi Ro Bán (0-100)", ascending=False).reset_index(
      drop=True
  )
 
  file_path = os.path.join(OUTPUT_DIR, "HOSE_Technical_Sell_Report.xlsx")
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Bộ Lọc Tín Hiệu Bán"
 
  ws.merge_cells("A1:J1")
  title_cell = ws["A1"]
  title_cell.value = (
      "BÁO CÁO TÍN HIỆU KỸ THUẬT BÁN & QUẢN TRỊ RỦI RO (HOSE WATCHLIST)"
  )
  title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
  title_cell.fill = PatternFill(
      start_color="C00000", end_color="C00000", fill_type="solid"
  )
  title_cell.alignment = Alignment(horizontal="center", vertical="center")
  ws.row_dimensions[1].height = 35
 
  headers = list(df.columns)
  ws.row_dimensions[3].height = 28
  header_fill = PatternFill(
      start_color="333333", end_color="333333", fill_type="solid"
  )
  header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
  border_thin = Border(
      left=Side(style="thin", color="D9D9D9"),
      right=Side(style="thin", color="D9D9D9"),
      top=Side(style="thin", color="D9D9D9"),
      bottom=Side(style="thin", color="D9D9D9"),
  )
 
  for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    cell.border = border_thin
 
  for row_idx, row_data in enumerate(df.to_dict("records"), 4):
    ws.row_dimensions[row_idx].height = 22
    for col_idx, key in enumerate(headers, 1):
      val = row_data[key]
      cell = ws.cell(row=row_idx, column=col_idx, value=val)
      cell.font = Font(name="Calibri", size=10)
      cell.border = border_thin
 
      if key in ["Mã CK", "Nhóm Ngành", "Chỉ báo RSI", "Mức Độ Ưu Tiên"]:
        cell.alignment = Alignment(horizontal="center", vertical="center")
      elif key == "Điểm Rủi Ro Bán (0-100)":
        cell.alignment = Alignment(horizontal="right", vertical="center")
      else:
        cell.alignment = Alignment(horizontal="left", vertical="center")
 
      if key == "Hành Động Khuyến Nghị":
        cell.font = Font(name="Calibri", size=10, bold=True)
        if "CẮT LỖ" in str(val) or "BÁN KHẨN CẤP" in str(val):
          cell.fill = PatternFill(
              start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="C00000"
          )
        elif "CHỐT LỜI" in str(val):
          cell.fill = PatternFill(
              start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="B28900"
          )
        elif "TIẾP TỤC NẮM GIỮ" in str(val):
          cell.fill = PatternFill(
              start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
          )
          cell.font = Font(
              name="Calibri", size=10, bold=True, color="375623"
          )
 
  for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
  ws.column_dimensions["J"].width = 38
 
  chart = BarChart()
  chart.type = "col"
  chart.style = 10
  chart.title = "Xếp Hạng Điểm Rủi Ro Kỹ Thuật Bán (Sell Risk Score)"
  chart.y_axis.title = "Điểm Rủi Ro (0-100)"
  chart.x_axis.title = "Mã Cổ Phiếu"
  data_ref = Reference(ws, min_col=3, min_row=3, max_row=3 + len(df))
  cats_ref = Reference(ws, min_col=1, min_row=4, max_row=3 + len(df))
  chart.add_data(data_ref, titles_from_data=True)
  chart.set_categories(cats_ref)
  chart.width = 18
  chart.height = 12
  ws.add_chart(chart, "L3")
 
  wb.save(file_path)
  print(f"   -> Đã lưu: {file_path}")
  return df
 
 
# ==========================================
# GỬI EMAIL
# ==========================================
def _top_table_html(df, columns, n=5):
  return df[columns].head(n).to_html(index=False, border=1, justify="left")
 
 
def build_email_body(df_watch, df_pullback, df_sell):
  now_vn = datetime.utcnow() + timedelta(hours=7)
  html = []
  html.append("<html><body style='font-family:Arial,sans-serif;font-size:14px;'>")
  html.append(
      f"<h2>HOSE Daily Screener - {now_vn.strftime('%d/%m/%Y %H:%M')} (GMT+7)</h2>"
  )
  if SIMULATED_DATA_NOTICE:
    html.append(
        f"<p style='color:#a00;'><b>{SIMULATED_DATA_NOTICE}</b></p>"
    )
 
  html.append("<h3>1. Top 5 Đánh giá tổng quan Watchlist</h3>")
  html.append(
      _top_table_html(
          df_watch, ["Mã CK", "Tổng Điểm Đánh Giá", "Khuyến Nghị Đánh Giá"]
      )
  )
  html.append("<h3>2. Top 5 Điểm mua Pullback / Test cung</h3>")
  html.append(
      _top_table_html(
          df_pullback, ["Mã CK", "Tổng Điểm Đánh Giá", "Chiến Lược Hành Động"]
      )
  )
  html.append("<h3>3. Top 5 Rủi ro bán cao nhất</h3>")
  html.append(
      _top_table_html(
          df_sell,
          ["Mã CK", "Điểm Rủi Ro Bán (0-100)", "Hành Động Khuyến Nghị"],
      )
  )
  html.append("<p>Chi tiết đầy đủ và biểu đồ xem trong 3 file Excel đính kèm.</p>")
  html.append("</body></html>")
  return "\n".join(html)
 
 
def send_email(subject, html_body, attachment_paths):
  sender = os.environ.get("MY_EMAIL", "").strip()
  # Mật khẩu ứng dụng Gmail hay bị dính dấu cách khi copy: loại bỏ mọi khoảng trắng
  password = "".join(os.environ.get("MAIL_PASSWORD", "").split())
  if not sender or not password:
    raise RuntimeError(
        "Thiếu biến môi trường MY_EMAIL hoặc MAIL_PASSWORD (kiểm tra GitHub Secrets)."
    )
 
  msg = MIMEMultipart()
  msg["From"] = sender
  msg["To"] = sender
  msg["Subject"] = subject
  msg.attach(MIMEText(html_body, "html", "utf-8"))
 
  for path in attachment_paths:
    with open(path, "rb") as f:
      part = MIMEBase("application", "octet-stream")
      part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header(
        "Content-Disposition",
        f'attachment; filename="{os.path.basename(path)}"',
    )
    msg.attach(part)
 
  raw_message = msg.as_string()
  last_error = None
 
  # Thử tối đa 3 vòng; mỗi vòng thử cổng 465 (SSL) rồi cổng 587 (STARTTLS)
  for attempt in range(1, 4):
    for mode, port in (("SSL-465", 465), ("STARTTLS-587", 587)):
      server = None
      try:
        if mode == "SSL-465":
          server = smtplib.SMTP_SSL("smtp.gmail.com", port, timeout=60)
        else:
          server = smtplib.SMTP("smtp.gmail.com", port, timeout=60)
          server.ehlo()
          server.starttls()
          server.ehlo()
        server.login(sender, password)
        server.sendmail(sender, [sender], raw_message)
        print(f"   -> Đã gửi email tới {sender} (vòng {attempt}, {mode})")
        return
      except smtplib.SMTPAuthenticationError as exc:
        raise RuntimeError(
            "Gmail từ chối đăng nhập: kiểm tra MY_EMAIL và mật khẩu ứng dụng "
            "(16 ký tự, không phải mật khẩu đăng nhập Gmail)."
        ) from exc
      except (smtplib.SMTPException, OSError) as exc:
        last_error = exc
        print(f"   -> Vòng {attempt}, {mode} lỗi: {exc!r}")
      finally:
        if server is not None:
          try:
            server.quit()
          except Exception:
            pass
    time.sleep(5)
 
  raise RuntimeError(f"Gửi email thất bại sau 3 vòng thử. Lỗi cuối: {last_error!r}")
 
 
# ==========================================
# HÀM CHÍNH ĐIỀU PHỐI TẤT CẢ
# ==========================================
if __name__ == "__main__":
  print("🎯 Bắt đầu chạy hệ thống tổng hợp toàn bộ báo cáo HOSE Watchlist...")
  df_watch = run_main_watchlist_report()
  df_pullback = run_pullback_analysis()
  df_sell = run_technical_sell_screener()
  print(f"✨ HOÀN TẤT! Cả 3 file Excel đã được cập nhật thành công tại {OUTPUT_DIR}")
 
  has_credentials = bool(
      os.environ.get("MY_EMAIL", "").strip()
      and os.environ.get("MAIL_PASSWORD", "").strip()
  )
  if has_credentials:
    print("📧 Đang gửi email báo cáo...")
    today_str = (datetime.utcnow() + timedelta(hours=7)).strftime("%d/%m/%Y")
    send_email(
        subject=f"HOSE Daily Screener {today_str}",
        html_body=build_email_body(df_watch, df_pullback, df_sell),
        attachment_paths=[
            os.path.join(OUTPUT_DIR, "HOSE_Watchlist_Report.xlsx"),
            os.path.join(OUTPUT_DIR, "HOSE_Watchlist_Pullback_Report.xlsx"),
            os.path.join(OUTPUT_DIR, "HOSE_Technical_Sell_Report.xlsx"),
        ],
    )
  elif os.environ.get("GITHUB_ACTIONS") == "true":
    raise RuntimeError(
        "Chạy trên GitHub Actions nhưng thiếu secrets MY_EMAIL / MAIL_PASSWORD."
    )
  else:
    print("ℹ️ Không có MY_EMAIL / MAIL_PASSWORD, bỏ qua bước gửi email (chạy cục bộ).")
 

         
